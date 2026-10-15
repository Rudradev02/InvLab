import io
from decimal import Decimal
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def export_to_excel(report_data: Dict[str, Any]) -> bytes:
    """
    Generate a styled Excel (.xlsx) file from report data dictionary.
    Uses openpyxl with InvLab brand styling.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(report_data.get("name", "Report"))[:31]

    # Brand Colors
    PRIMARY_COLOR = "173D32"  # InvLab Forest Green
    ZEBRA_COLOR = "F8F9F8"    # Subtle light ivory/green
    BORDER_COLOR = "E3E1DA"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10)
    zebra_fill = PatternFill(start_color=ZEBRA_COLOR, end_color=ZEBRA_COLOR, fill_type="solid")

    thin_side = Side(border_style="thin", color=BORDER_COLOR)
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    columns = report_data.get("columns", [])
    rows = report_data.get("rows", [])

    # Write Headers (Row 1)
    ws.row_dimensions[1].height = 26
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col.get("label", ""))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = cell_border

    # Write Data Rows (Row 2 onwards)
    for row_idx, row in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)

        for col_idx, col in enumerate(columns, start=1):
            key = col.get("key", "")
            val = row.get(key, "")

            # Convert Decimal to float for Excel numeric formatting
            if isinstance(val, Decimal):
                val = float(val)

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = cell_border

            if is_even:
                cell.fill = zebra_fill

            # Alignment and formatting
            align = col.get("align", "left")
            if align == "right" or col.get("is_currency", False) or isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if col.get("is_currency", False) or isinstance(val, float):
                    cell.number_format = "₹#,##0.00"
                elif isinstance(val, int):
                    cell.number_format = "#,##0"
            elif align == "center":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col_idx, col in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col.get("label", "")))
        for row in rows[:200]:  # sample up to 200 rows for performance
            val_str = str(row.get(col.get("key", ""), "") or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 45))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
